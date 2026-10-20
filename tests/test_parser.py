import pytest
import pymupdf
import docx
import openpyxl
import shutil
from PIL import Image, ImageDraw
from app.parsers.base import ParserError
from app.parsers.json_parser import JsonParser
from pathlib import Path
from app.parsers.text import TextParser
from app.parsers.csv_parser import CsvParser
from app.parsers.code import CodeParser
from app.parsers.pdf import PdfParser
from app.parsers.docx import DocxParser
from app.parsers.xlsx import XlsxParser
from app.parsers.image_parser import ImageParser

def test_text_parser_can_handle():
    parser = TextParser()
    assert parser.can_handle("notes.txt") is True
    assert parser.can_handle("README.md") is True
    assert parser.can_handle("document.pdf") is False


def test_text_parser_parse_utf8(tmp_path: Path):
    parser = TextParser()
    sample_file = tmp_path / "hello.txt"
    sample_file.write_text("Hello, File2MD!\nBaris kedua.", encoding="utf-8")

    result = parser.parse(sample_file)
    assert result.content == "Hello, File2MD!\nBaris kedua."
    assert result.metadata["encoding"] == "utf-8"
    assert result.metadata["line_count"] == 2
    assert len(result.warnings) == 0


def test_text_parser_parse_latin1(tmp_path: Path):
    parser = TextParser()
    sample_file = tmp_path / "iso.txt"
    # Tulis byte latin-1 (karakter 'é' dalam latin-1 adalah b'\xe9', bukan valid UTF-8)
    sample_file.write_bytes(b"Caf\xe9")

    result = parser.parse(sample_file)
    assert "Café" in result.content
    assert len(result.warnings) > 0

def test_json_parser_can_handle():
    parser = JsonParser()
    assert parser.can_handle("data.json") is True
    assert parser.can_handle("data.txt") is False

def test_json_parser_valid(tmp_path: Path):
    parser = JsonParser()
    sample = tmp_path / "valid.json"
    sample.write_text('{"name": "File2MD", "version": 1}', encoding="utf-8")
    result = parser.parse(sample)
    assert result.content.startswith("```json\n")
    assert result.content.endswith("\n```")
    assert '"name": "File2MD"' in result.content
    assert result.metadata["root_type"] == "dict"
    assert result.metadata["item_count"] == 2
    
def test_json_parser_invalid(tmp_path: Path):
    parser = JsonParser()
    sample = tmp_path / "invalid.json"
    sample.write_text('{"name": "File2MD", UNQUOTED}', encoding="utf-8")
    with pytest.raises(ParserError) as exc_info:
        parser.parse(sample)
    assert "Format JSON tidak valid" in str(exc_info.value)

def test_csv_parser_can_handle():
    parser = CsvParser()
    assert parser.can_handle("data.csv") is True
    assert parser.can_handle("data.tsv") is True
    assert parser.can_handle("data.json") is False

def test_csv_parser_valid(tmp_path: Path):
    parser = CsvParser()
    sample = tmp_path / "test.csv"
    sample.write_text("Nama,Usia,Kota\nBudi,25,Jakarta\nSiti,30,Surabaya\n", encoding="utf-8")
    result = parser.parse(sample)
    assert "| Nama | Usia | Kota |" in result.content
    assert "| --- | --- | --- |" in result.content
    assert "| Budi | 25 | Jakarta |" in result.content
    assert result.metadata["row_count"] == 3
    assert result.metadata["col_count"] == 3

def test_csv_parser_pipe_escaping(tmp_path: Path):
    parser = CsvParser()
    sample = tmp_path / "pipe.csv"
    sample.write_text('Produk,Deskripsi\nLaptop,"Model A | 16GB RAM"\n', encoding="utf-8")
    result = parser.parse(sample)
    # Memastikan pipe di dalam data di-escape agar tidak merusak tabel
    assert r"Model A \| 16GB RAM" in result.content
    
def test_csv_parser_semicolon_delimiter(tmp_path: Path):
    parser = CsvParser()
    sample = tmp_path / "semicolon.csv"
    sample.write_text("ID;Barang;Harga\n1;Buku;50000\n", encoding="utf-8")
    result = parser.parse(sample)
    assert result.metadata["delimiter"] == ";"
    assert "| ID | Barang | Harga |" in result.content

def test_code_parser_can_handle():
    parser = CodeParser()
    assert parser.can_handle("script.py") is True
    assert parser.can_handle("app.ts") is True
    assert parser.can_handle("Dockerfile") is True
    assert parser.can_handle("report.pdf") is False

def test_code_parser_parse_python(tmp_path: Path):
    parser = CodeParser()
    sample = tmp_path / "main.py"
    sample.write_text("def hello():\n    print('world')\n", encoding="utf-8")
    result = parser.parse(sample)
    assert result.content.startswith("```python\n")
    assert result.content.endswith("\n```")
    assert result.metadata["language"] == "python"
    assert result.metadata["line_count"] == 3

def test_code_parser_safe_fencing(tmp_path: Path):
    parser = CodeParser()
    sample = tmp_path / "snippet.js"
    # Kode JS yang mengandung triple backtick di dalamnya
    sample.write_text("const md = '```markdown';\n", encoding="utf-8")
    result = parser.parse(sample)
    # Pagar luar harus menggunakan minimal 4 backtick agar tidak bentrok
    assert result.content.startswith("````javascript\n")
    assert result.content.endswith("\n````")

def test_pdf_parser_can_handle():
    parser = PdfParser()
    assert parser.can_handle("document.pdf") is True
    assert parser.can_handle("document.txt") is False

def test_pdf_parser_text_extraction(tmp_path: Path):
    parser = PdfParser()
    pdf_path = tmp_path / "sample.pdf"
    # Buat PDF dummy 2 halaman dengan PyMuPDF
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Halo Halaman Satu!")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Halo Halaman Dua!")
    doc.save(str(pdf_path))
    doc.close()
    result = parser.parse(pdf_path)
    assert "<!-- Page 1 -->" in result.content
    assert "Halo Halaman Satu!" in result.content
    assert "<!-- Page 2 -->" in result.content
    assert "Halo Halaman Dua!" in result.content
    assert result.metadata["page_count"] == 2
    assert result.metadata["text_pages"] == 2

def test_docx_parser_can_handle():
    parser = DocxParser()
    assert parser.can_handle("laporan.docx") is True
    assert parser.can_handle("laporan.pdf") is False

def test_docx_parser_structure(tmp_path: Path):
    parser = DocxParser()
    docx_path = tmp_path / "sample.docx"
    # Buat dokumen Word secara programmatic untuk pengujian
    doc = docx.Document()
    doc.add_heading("Bab 1 Pendahuluan", level=1)
    doc.add_paragraph("Ini adalah isi paragraf pengantar.")
    doc.add_paragraph("Poin A", style="List Bullet")
    doc.add_paragraph("Poin B", style="List Bullet")
    # Tambahkan tabel
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Header 1"
    tbl.rows[0].cells[1].text = "Header 2"
    tbl.rows[1].cells[0].text = "Data 1"
    tbl.rows[1].cells[1].text = "Data 2 | Spesial"
    doc.save(str(docx_path))
    result = parser.parse(docx_path)
    assert "# Bab 1 Pendahuluan" in result.content
    assert "Ini adalah isi paragraf pengantar." in result.content
    assert "- Poin A" in result.content
    assert "| Header 1 | Header 2 |" in result.content
    assert r"Data 2 \| Spesial" in result.content
    assert result.metadata["paragraph_count"] >= 3
    assert result.metadata["table_count"] == 1

def test_xlsx_parser_can_handle():
    parser = XlsxParser()
    assert parser.can_handle("data.xlsx") is True
    assert parser.can_handle("macro.xlsm") is True
    assert parser.can_handle("data.csv") is False
    
def test_xlsx_parser_multi_sheets(tmp_path: Path):
    parser = XlsxParser()
    xlsx_path = tmp_path / "test.xlsx"
    # Buat file Excel dengan 2 sheet
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Penjualan"
    ws1.append(["Barang", "Qty", "Harga"])
    ws1.append(["Buku", 10, 50000])
    ws1.append(["Pulpen | Biru", 2, 10000])
    ws2 = wb.create_sheet(title="Info")
    ws2.append(["Catatan"])
    ws2.append(["Data bulan Oktober"])
    wb.save(xlsx_path)
    wb.close()
    result = parser.parse(xlsx_path)
    assert "## Sheet: Penjualan" in result.content
    assert "| Barang | Qty | Harga |" in result.content
    assert "| Buku | 10 | 50000 |" in result.content
    assert r"Pulpen \| Biru" in result.content
    assert "## Sheet: Info" in result.content
    assert "| Catatan |" in result.content
    assert result.metadata["total_sheets"] == 2

def test_image_parser_can_handle():
    parser = ImageParser()
    assert parser.can_handle("foto.png") is True
    assert parser.can_handle("banner.jpg") is True
    assert parser.can_handle("icon.webp") is True
    assert parser.can_handle("document.pdf") is False

def test_image_parser_metadata_extraction(tmp_path: Path):
    parser = ImageParser()
    img_path = tmp_path / "test.png"
    # Buat gambar 800x600 px (Rasio 4:3)
    img = Image.new("RGB", (800, 600), color="blue")
    img.save(img_path, format="PNG")
    result = parser.parse(img_path)
    assert "| **Format** | PNG |" in result.content
    assert "| **Dimensi** | 800 × 600 px |" in result.content
    assert "| **Rasio Aspek** | 4:3 |" in result.content
    assert "| **Mode Warna** | RGB |" in result.content
    assert "data:image" not in result.content, "Hindari base64 bloat agar aman untuk chat AI"
    assert result.metadata["width"] == 800
    assert result.metadata["height"] == 600
    assert result.metadata["aspect_ratio"] == "4:3"

def test_image_parser_ocr_integration(tmp_path: Path):
    # Tes ini otomatis berjalan jika tesseract sudah diinstall di sistem
    if not shutil.which("tesseract"):
        pytest.skip("Tesseract tidak terpasang di sistem.")
    parser = ImageParser()
    img_path = tmp_path / "ocr_test.png"
    # Buat gambar putih bersih dengan teks hitam tebal
    img = Image.new("RGB", (600, 200), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((30, 80), "FILE2MD OCR TEST", fill="black")
    img.save(img_path, format="PNG")
    result = parser.parse(img_path)
    assert result.metadata["has_ocr_text"] is True
    assert "FILE2MD" in result.content.upper()
    assert "Ekstraksi Teks (OCR Lokal)" in result.content