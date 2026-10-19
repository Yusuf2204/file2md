from pathlib import Path
import openpyxl
from app.parsers.base import BaseParser, ParseResult, ParserError


class XlsxParser(BaseParser):
    SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

    def can_handle(self, filename: str, mime_type: str | None = None) -> bool:
        ext = Path(filename).suffix.lower()
        if ext in self.SUPPORTED_EXTENSIONS:
            return True
        if mime_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        }:
            return True
        return False

    def _sanitize_cell(self, value) -> str:
        if value is None:
            return ""
        text = str(value).strip().replace("|", r"\|")
        return text.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")

    def _convert_sheet(self, ws) -> str | None:
        # Baca semua baris sebagai tuple nilai
        raw_rows = list(ws.iter_rows(values_only=True))

        # Filter baris yang benar-benar memiliki isi
        non_empty_rows: list[list[str]] = []
        max_cols = 0

        for row in raw_rows:
            if not row:
                continue
            # Periksa apakah ada sel yang bukan None / bukan whitespace
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                cleaned_row = [self._sanitize_cell(c) for c in row]
                # Pangkas trailing kolom kosong di kanan
                while cleaned_row and cleaned_row[-1] == "":
                    cleaned_row.pop()
                if cleaned_row:
                    non_empty_rows.append(cleaned_row)
                    max_cols = max(max_cols, len(cleaned_row))

        if not non_empty_rows or max_cols == 0:
            return "*(Sheet kosong)*"

        # Baris pertama sebagai header
        header_row = non_empty_rows[0]
        # Pad jika header lebih pendek dari max_cols
        while len(header_row) < max_cols:
            header_row.append(f"Col_{len(header_row)+1}")

        # Ganti header kosong jika ada
        headers = [h or f"Col_{i+1}" for i, h in enumerate(header_row[:max_cols])]
        separator = ["---"] * max_cols

        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(separator) + " |",
        ]

        # Baris data
        for row in non_empty_rows[1:]:
            cells = row[:max_cols]
            while len(cells) < max_cols:
                cells.append("")
            lines.append("| " + " | ".join(cells) + " |")

        return "\n".join(lines)

    def parse(self, file_path: Path) -> ParseResult:
        warnings: list[str] = []

        try:
            wb = openpyxl.load_workbook(filename=file_path, data_only=True, read_only=True)
        except Exception as exc:
            raise ParserError(f"Gagal membaca file Excel (XLSX): {str(exc)}") from exc

        sheet_sections: list[str] = []
        total_rows = 0

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                table_md = self._convert_sheet(ws)
                section = f"## Sheet: {sheet_name}\n\n{table_md}"
                sheet_sections.append(section)
        finally:
            wb.close()

        if not sheet_sections:
            warnings.append("Buku kerja Excel tidak memiliki sheet.")

        markdown_output = "\n\n---\n\n".join(sheet_sections)

        return ParseResult(
            content=markdown_output,
            metadata={
                "parser": "XlsxParser",
                "sheet_names": wb.sheetnames,
                "total_sheets": len(wb.sheetnames),
            },
            warnings=warnings,
        )